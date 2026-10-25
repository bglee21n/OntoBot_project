"""사용자 매뉴얼 빌드 — docs/manual/user/*.md(번호 순)를 Word(.docx)와 PDF 로 묶는다.

    uv run --with python-docx --with markdown --with playwright python docs/manual/tools/build_manual.py

산출물은 ``dist/manual/``(git 이 무시한다). 원본은 Markdown 이 정본이다 — Word 를 직접 고치지 말 것.
Markdown 은 이 매뉴얼이 쓰는 만큼만 읽는다: 제목(#~####), 문단, 굵게(**), 코드(`), 목록(-, 1.), 표(|),
그림(![설명](경로)), 인용(>) — 알림 상자로, 코드 블록(```), 쪽 나눔(<!-- pagebreak -->).

Word 의 목차는 필드라서 처음 열 때 "필드 업데이트"를 물으면 예를 누른다(쪽 번호가 채워진다).
PDF 는 같은 원본을 HTML 로 바꿔 Edge 로 인쇄한다 — 쪽수 확인에도 쓴다.
"""

from __future__ import annotations

import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

MANUAL = Path(__file__).resolve().parents[1]
SRC = MANUAL / "user"
ROOT = MANUAL.parents[1]
OUT = ROOT / "dist" / "manual"
TITLE = "OntoBot 사용자 매뉴얼"
FONT = "맑은 고딕"
MONO = "Consolas"
PAGE_W = Cm(21.0)
MARGIN = Cm(2.0)
BODY_W = PAGE_W - 2 * MARGIN

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def chapters() -> list[Path]:
    return sorted(p for p in SRC.glob("*.md") if re.match(r"^\d\d", p.name))


# ---------------------------------------------------------------- Word


def set_font(run, size: float | None = None, bold: bool | None = None, mono: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = MONO if mono else FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_inline(par, text: str, size: float | None = None, base_bold: bool = False) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            set_font(r, size, True)
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            set_font(r, size, base_bold, mono=True, color=RGBColor(0x9F, 0x12, 0x39))
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            r = par.add_run(m.group(1) if m else part)
            set_font(r, size, base_bold, color=RGBColor(0x1D, 0x4E, 0xD8))
        else:
            r = par.add_run(part)
            set_font(r, size, base_bold)


def shade(cell_or_par, hex_fill: str) -> None:
    el = cell_or_par._tc if hasattr(cell_or_par, "_tc") else cell_or_par._p
    pr = el.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def border_par(par, color: str = "2563EB", side: str = "left", size: int = 18) -> None:
    ppr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement(f"w:{side}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(size))
    b.set(qn("w:space"), "6")
    b.set(qn("w:color"), color)
    bdr.append(b)
    ppr.append(bdr)


def field(par, instr: str) -> None:
    r = par.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    r._r.append(b)
    r2 = par.add_run()
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = instr
    r2._r.append(t)
    r3 = par.add_run()
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    r3._r.append(s)
    r4 = par.add_run("(목차는 Word 에서 열 때 필드 업데이트로 채워집니다)" if instr.startswith("TOC") else "1")
    set_font(r4, 10)
    r5 = par.add_run()
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r5._r.append(e)


def setup(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = PAGE_W, Cm(29.7)
    sec.left_margin = sec.right_margin = MARGIN
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.line_spacing = 1.35
    for lvl, size, color in [(1, 20, "1E3A8A"), (2, 15, "1E40AF"), (3, 12.5, "1F2937"), (4, 11, "374151")]:
        hs = doc.styles[f"Heading {lvl}"]
        hs.font.name = FONT
        hs.font.size = Pt(size)
        hs.font.bold = True
        hs.font.color.rgb = RGBColor.from_string(color)
        rpr = hs.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            fonts.set(qn(a), FONT)
        hs.paragraph_format.space_before = Pt(18 if lvl == 1 else 12)
        hs.paragraph_format.space_after = Pt(8)
        hs.paragraph_format.keep_with_next = True


def cover(doc: Document) -> None:
    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("OntoBot")
    set_font(r, 40, True, color=RGBColor(0x25, 0x63, 0xEB))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("사용자 매뉴얼")
    set_font(r, 28, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("실제 화면으로 따라 하는 온톨로지 기반 챗봇 사용법")
    set_font(r, 14, color=RGBColor(0x4B, 0x55, 0x63))
    for _ in range(10):
        doc.add_paragraph()
    for line in ("일반 사용자 · 승인자 · 현업 전문가 · 모델러 · 챗봇 빌더 · 관리자", f"{date.today():%Y. %m. %d.} 판"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        set_font(r, 11, color=RGBColor(0x4B, 0x55, 0x63))
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h = doc.add_paragraph()
    r = h.add_run("차례")
    set_font(r, 20, True, color=RGBColor(0x1E, 0x3A, 0x8A))
    field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def footer(doc: Document) -> None:
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("— ")
        set_font(r, 9)
        field(p, "PAGE")
        r = p.add_run(" —")
        set_font(r, 9)
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run(TITLE)
        set_font(r, 8.5, color=RGBColor(0x6B, 0x72, 0x80))


class WordWriter:
    def __init__(self) -> None:
        self.doc = Document()
        setup(self.doc)
        self.figure = 0
        self.chapter_no = 0

    def heading(self, text: str, level: int) -> None:
        if level == 1:
            self.chapter_no += 1
            if self.chapter_no > 1:
                self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        self.doc.add_heading(text, level=level)

    def paragraph(self, text: str) -> None:
        add_inline(self.doc.add_paragraph(), text)

    def bullets(self, items: list[tuple[int, bool, str]]) -> None:
        # 번호 목록은 번호를 직접 붙인다(Word 번호 매기기는 문서 전체에서 이어져 매 목록 1 부터 시작하지 않는다).
        counters = [0, 0]
        for indent, ordered, text in items:
            if ordered:
                counters[indent] += 1
                if indent == 0:
                    counters[1] = 0
                p = self.doc.add_paragraph()
                fmt = p.paragraph_format
                fmt.left_indent = Cm(0.9 + 0.8 * indent)
                fmt.first_line_indent = Cm(-0.6)
                fmt.space_after = Pt(2)
                r = p.add_run(f"{counters[indent]}. ")
                set_font(r, None, True, color=RGBColor(0x25, 0x63, 0xEB))
            else:
                p = self.doc.add_paragraph(style="List Bullet 2" if indent else "List Bullet")
                p.paragraph_format.space_after = Pt(2)
            add_inline(p, text)

    def table(self, rows: list[list[str]]) -> None:
        head, body = rows[0], rows[1:]
        n = len(head)
        t = self.doc.add_table(rows=1 + len(body), cols=n)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = col_widths(rows)
        for i, row in enumerate([head, *body]):
            for j in range(n):
                cell = t.cell(i, j)
                cell.width = widths[j]
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                add_inline(p, row[j] if j < len(row) else "", 9.5, base_bold=(i == 0))
                if i == 0:
                    shade(cell, "E8EEFC")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def image(self, path: Path, caption: str) -> None:
        from PIL import Image

        with Image.open(path) as im:
            w, h = im.size
        width = min(BODY_W, Cm(16.5) * min(1.0, w / 1280))
        max_h = Cm(21.0)
        if width * h / w > max_h:
            width = int(max_h * w / h)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=width)
        self.figure += 1
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(f"그림 {self.figure}. {caption}")
        set_font(r, 9, color=RGBColor(0x4B, 0x55, 0x63))

    def note(self, lines: list[str]) -> None:
        kind = "참고"
        first = lines[0] if lines else ""
        m = re.match(r"\*\*(주의|중요|참고|따라 하기|요령|확인|해 보기|정답|알아 두기)[^*]*\*\*", first)
        color, fill = "2563EB", "EFF6FF"
        if m:
            kind = m.group(1)
            if kind in ("주의", "중요"):
                color, fill = "DC2626", "FEF2F2"
            elif kind in ("요령", "따라 하기", "해 보기"):
                color, fill = "059669", "ECFDF5"
        for line in lines:
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.3)
            p.paragraph_format.space_after = Pt(1)
            border_par(p, color)
            shade(p, fill)
            add_inline(p, line)

    def code(self, lines: list[str]) -> None:
        for line in lines:
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            shade(p, "F3F4F6")
            r = p.add_run(line if line else " ")
            set_font(r, 9, mono=True)

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path: Path) -> None:
        footer(self.doc)
        self.doc.save(str(path))


def col_widths(rows: list[list[str]]):
    n = len(rows[0])
    lens = [max(len(r[j]) if j < len(r) else 0 for r in rows) for j in range(n)]
    lens = [min(max(x, 4), 60) for x in lens]
    total = sum(lens)
    return [int(BODY_W * x / total) for x in lens]


# ---------------------------------------------------------------- Markdown → 블록


def parse(md: str, base: Path, w: WordWriter) -> None:
    lines = md.splitlines()
    i = 0
    para: list[str] = []

    def flush() -> None:
        if para:
            w.paragraph(" ".join(s.strip() for s in para))
            para.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            flush()
            i += 1
            continue
        if s == "<!-- pagebreak -->":
            flush()
            w.page_break()
            i += 1
            continue
        if s.startswith("<!--"):
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            flush()
            w.heading(m.group(2), len(m.group(1)))
            i += 1
            continue
        if s.startswith("```"):
            flush()
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            w.code(block)
            i += 1
            continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", s)
        if m:
            flush()
            w.image((base / m.group(2)).resolve(), m.group(1))
            i += 1
            continue
        if s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            w.table(rows)
            continue
        if s.startswith(">"):
            flush()
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip()[1:].strip())
                i += 1
            merged: list[str] = []
            for b in block:
                if not b:
                    continue
                if merged and not b.startswith(("- ", "**")) and not re.match(r"^\d+\.", b):
                    merged[-1] += " " + b
                else:
                    merged.append(b)
            w.note(merged)
            continue
        if re.match(r"^(\s*)([-*]|\d+\.)\s+", line):
            flush()
            items = []
            while i < len(lines) and (re.match(r"^(\s*)([-*]|\d+\.)\s+", lines[i]) or (lines[i].startswith("   ") and lines[i].strip() and items)):
                m2 = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", lines[i])
                if m2:
                    items.append((1 if len(m2.group(1)) >= 2 else 0, m2.group(2)[0].isdigit(), m2.group(3)))
                else:
                    ind, o, t = items[-1]
                    items[-1] = (ind, o, t + " " + lines[i].strip())
                i += 1
            w.bullets(items)
            continue
        if s == "---":
            flush()
            i += 1
            continue
        para.append(s)
        i += 1
    flush()


# ---------------------------------------------------------------- HTML → PDF


CSS = """
@page { size: A4; margin: 20mm 18mm 18mm 18mm; }
body { font-family: 'Malgun Gothic', sans-serif; font-size: 10.5pt; line-height: 1.6; color: #111827; }
h1 { page-break-before: always; color: #1e3a8a; font-size: 20pt; border-bottom: 2px solid #2563eb; padding-bottom: 4px; }
h1.first, .cover + h1 { page-break-before: always; }
h2 { color: #1e40af; font-size: 15pt; margin-top: 18px; }
h3 { font-size: 12.5pt; margin-top: 14px; }
h4 { font-size: 11pt; }
h2, h3, h4 { page-break-after: avoid; }
table { border-collapse: collapse; width: 100%; font-size: 9.5pt; margin: 6px 0 10px; page-break-inside: auto; }
th, td { border: 1px solid #cbd5e1; padding: 3px 6px; vertical-align: top; }
th { background: #e8eefc; }
tr { page-break-inside: avoid; }
code { font-family: Consolas, monospace; color: #9f1239; font-size: 9.5pt; }
pre { background: #f3f4f6; padding: 8px; font-size: 9pt; white-space: pre-wrap; }
pre code { color: #111827; }
blockquote { border-left: 4px solid #2563eb; background: #eff6ff; margin: 8px 0; padding: 6px 10px; }
figure { text-align: center; margin: 10px 0; page-break-inside: avoid; }
figure img { max-width: 100%; max-height: 210mm; border: 1px solid #e5e7eb; }
figcaption { font-size: 9pt; color: #4b5563; margin-top: 3px; }
.cover { text-align: center; padding-top: 70mm; page-break-after: always; }
.cover h1 { page-break-before: avoid; border: none; font-size: 40pt; color: #2563eb; }
.pb { page-break-after: always; }
"""


def build_html(paths: list[Path]) -> str:
    import markdown

    parts = [
        f"<div class='cover'><h1>OntoBot</h1><div style='font-size:28pt;font-weight:bold'>사용자 매뉴얼</div>"
        f"<p style='font-size:14pt;color:#4b5563'>실제 화면으로 따라 하는 온톨로지 기반 챗봇 사용법</p>"
        f"<p style='margin-top:80mm;color:#4b5563'>{date.today():%Y. %m. %d.} 판</p></div>"
    ]
    fig = 0
    for p in paths:
        text = p.read_text(encoding="utf-8").replace("<!-- pagebreak -->", "<div class='pb'></div>")

        def img(m: re.Match[str]) -> str:
            nonlocal fig
            fig += 1
            src = (p.parent / m.group(2)).resolve().as_uri()
            return f"\n<figure><img src='{src}'><figcaption>그림 {fig}. {m.group(1)}</figcaption></figure>\n"

        text = re.sub(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", img, text, flags=re.M)
        parts.append(markdown.markdown(text, extensions=["tables", "fenced_code"]))
    return f"<!doctype html><html lang='ko'><head><meta charset='utf-8'><style>{CSS}</style></head><body>{''.join(parts)}</body></html>"


def build_pdf(html: str, out: Path) -> int:
    from playwright.sync_api import sync_playwright

    tmp = out.with_suffix(".html")
    tmp.write_text(html, encoding="utf-8")
    with sync_playwright() as p:
        b = p.chromium.launch(channel="msedge")
        pg = b.new_page()
        pg.goto(tmp.as_uri())
        pg.wait_for_load_state("networkidle")
        pg.pdf(
            path=str(out),
            format="A4",
            print_background=True,
            display_header_footer=True,
            header_template=f"<div style='font-size:7px;width:100%;text-align:right;padding-right:18mm;color:#6b7280'>{TITLE}</div>",
            footer_template="<div style='font-size:8px;width:100%;text-align:center;color:#374151'>— <span class='pageNumber'></span> / <span class='totalPages'></span> —</div>",
            margin={"top": "20mm", "bottom": "18mm", "left": "18mm", "right": "18mm"},
        )
        b.close()
    data = out.read_bytes()
    return len(re.findall(rb"/Type\s*/Page[^s]", data))


def main() -> None:
    paths = chapters()
    if not paths:
        sys.exit(f"원본이 없습니다: {SRC}")
    OUT.mkdir(parents=True, exist_ok=True)
    w = WordWriter()
    cover(w.doc)
    for p in paths:
        parse(p.read_text(encoding="utf-8"), p.parent, w)
    docx_path = OUT / "OntoBot-사용자-매뉴얼.docx"
    w.save(docx_path)
    print(f"Word: {docx_path} (그림 {w.figure}장)")
    if "--no-pdf" not in sys.argv:
        pdf_path = OUT / "OntoBot-사용자-매뉴얼.pdf"
        pages = build_pdf(build_html(paths), pdf_path)
        print(f"PDF : {pdf_path} ({pages}쪽)")


if __name__ == "__main__":
    main()
